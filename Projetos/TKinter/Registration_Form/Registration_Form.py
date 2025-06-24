from tkinter import *
from openpyxl import load_workbook  

wb = load_workbook('C:Users/augus/Documents/Projetos VSCode/Projetos Python/TKinter/Projetos/Registration_Form/Registration_Form.xlsx')
ws = wb.active

def init_excel():
    headers = ["Name", "Course", "Semester", "Form No.", "Contact No.", "Email", "Address"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i).value = h
    wb.save('C:/Users/augus/Documents/Projetos VSCode/Projetos Python/TKinter/Projetos/Registration_Form/Registration_Form.xlsx')

def insert_data():
    if all(f.get() for f in entries):
        row = ws.max_row + 1
        for i, f in enumerate(entries,1):
            ws.cell(row=row, column=i).value = f.get()
        wb.save('C:/Users/augus/Documents/Projetos VSCode/Projetos Python/TKinter/Projetos/Registration_Form/Registration_Form.xlsx')
        clear_fields()
    else:
        print("Please fill all fields")

def clear_fields():
    for f in entries:
        f.delete(0, END)
        
def focus_text(entry):
    entry.focus_set()
    
init_excel()
root = Tk()
root.title("Registration Form")
root.geometry("500x300")

labels = ["Name", "Course", "Semester", "Form No.", "Contact No.", "Email", "Address"]
entries = [Entry(root) for _ in labels]

for i, lbl in enumerate(labels):
    Label(root,text = lbl).grid(row=i+1, column = 0)
    entries[i].grid(row=i+1, column = 1, ipadx=100)
    if i < len(labels) - 1:
        entries[i].bind("<Return>", lambda e, nf = entries[i+1]: focus_text(nf))
Button(root, text="Submit", fg = "black", bg = "white", command=insert_data).grid(row=len(labels)+1, column=1)

root.mainloop()